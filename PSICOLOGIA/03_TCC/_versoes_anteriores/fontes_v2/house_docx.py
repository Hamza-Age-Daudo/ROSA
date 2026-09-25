"""Construtor de documentos Word no estilo da casa UniLurio (FCS).

Aplica as regras fixas dos trabalhos da pasta UNILURIO: A4, margens 3/3/2,5/2,5 cm,
Times New Roman 12, espacamento 1,5, 6 pts depois do paragrafo, texto justificado,
numeracao no canto inferior direito, titulos ate ao 2.o nivel a negrito, tabelas de
tres linhas sem grelhas verticais, ausencia de travessoes e ortografia anterior ao
Acordo Ortografico de 1990.

Copia local da versao de _FERRAMENTAS, alargada com: cabecalho e rodape a 1,5 cm,
indice automatico, seccoes com numeracao propria (romana nos pre-textuais), formulas
com expoentes reais do Word e referencias com avanco de 0,75 cm.
"""
from __future__ import annotations

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONTE = "Times New Roman"
CORPO = Pt(12)
PRETO = RGBColor(0, 0, 0)
ESPACAMENTO = 1.5
DEPOIS = Pt(6)
DISTANCIA_CABECALHO_RODAPE = Cm(1.5)
AVANCO_REFERENCIAS = Cm(0.75)
SUCESSORES_PGNUMTYPE = ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
                        "w:printerSettings", "w:sectPrChange")

TRACOS = {"—": ", ", "–": "-", "−": "-", "→": " para ", "≈": "cerca de "}

# Correccoes pos-Acordo para pre-Acordo de 1990, aplicadas a todo o texto do documento.
PRE_1990 = [
    (r"\bprotecao\b", "proteccao"), (r"\bobjetiv", "objectiv"), (r"\batividad", "actividad"),
    (r"\bfatores\b", "factores"), (r"\bfator\b", "factor"), (r"\bsetor", "sector"),
    (r"\bcorret", "correct"), (r"\baspeto", "aspecto"), (r"\bafet", "afect"),
    (r"\bcontato", "contacto"), (r"\bespet", "espect"), (r"\bperspetiv", "perspectiv"),
    (r"\bad(o)ção\b", "adopção"), (r"\bproteção\b", "protecção"),
    (r"\bdeteção\b", "detecção"), (r"\bseleção\b", "selecção"),
    (r"\bdireção\b", "direcção"), (r"\breação\b", "reacção"),
    (r"\binteração\b", "interacção"), (r"\bcorreção\b", "correcção"),
    (r"\binfeção\b", "infecção"), (r"\bexceção\b", "excepção"),
]


def limpar(texto: str) -> str:
    """Remove travessoes e simbolos proibidos e repoe a ortografia anterior a 1990."""
    for mau, bom in TRACOS.items():
        texto = texto.replace(mau, bom)
    for padrao, certo in PRE_1990:
        texto = re.sub(padrao, certo, texto)
    return texto


def _formatar_run(run, tamanho=CORPO, negrito=False, italico=False):
    run.font.name = FONTE
    run.font.size = tamanho
    run.font.bold = negrito
    run.font.italic = italico
    run.font.color.rgb = PRETO
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONTE)


def _formatar_paragrafo(par, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, espaco=ESPACAMENTO,
                        depois=DEPOIS, antes=Pt(0)):
    pf = par.paragraph_format
    pf.alignment = alinhamento
    pf.line_spacing = espaco
    pf.space_after = depois
    pf.space_before = antes


def novo_documento() -> Document:
    """Devolve um documento vazio com a pagina, os estilos e o rodape da casa."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(3)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header_distance = sec.footer_distance = DISTANCIA_CABECALHO_RODAPE

    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = CORPO
    normal.font.color.rgb = PRETO
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
    normal.paragraph_format.line_spacing = ESPACAMENTO
    normal.paragraph_format.space_after = DEPOIS
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for nivel, tamanho in ((1, Pt(13)), (2, CORPO), (3, CORPO), (4, CORPO)):
        estilo = doc.styles[f"Heading {nivel}"]
        estilo.font.name = FONTE
        estilo.font.size = tamanho
        estilo.font.bold = nivel <= 2
        estilo.font.italic = False
        estilo.font.color.rgb = PRETO
    return doc


def numerar_paginas(doc: Document, ocultar_na_capa: bool = True) -> None:
    """Coloca o numero de pagina no canto inferior direito, sem mais nada no rodape.

    Por omissao a capa nao mostra o numero, embora conte para a numeracao seguinte. So a
    primeira pagina da primeira seccao fica sem numero; as seccoes seguintes mostram-no
    logo na primeira pagina.
    """
    for indice, sec in enumerate(doc.sections):
        esconder = ocultar_na_capa and indice == 0
        sec.different_first_page_header_footer = esconder
        if esconder:
            primeiro = sec.first_page_footer
            primeiro.is_linked_to_previous = False
            for par in primeiro.paragraphs:
                par.text = ""
        rodape = sec.footer
        rodape.is_linked_to_previous = False
        par = rodape.paragraphs[0] if rodape.paragraphs else rodape.add_paragraph()
        par.text = ""
        _formatar_paragrafo(par, WD_ALIGN_PARAGRAPH.RIGHT, 1.0, Pt(0))
        run = par.add_run()
        _formatar_run(run)
        inicio = OxmlElement("w:fldChar")
        inicio.set(qn("w:fldCharType"), "begin")
        instrucao = OxmlElement("w:instrText")
        instrucao.set(qn("xml:space"), "preserve")
        instrucao.text = " PAGE "
        fim = OxmlElement("w:fldChar")
        fim.set(qn("w:fldCharType"), "end")
        for elem in (inicio, instrucao, fim):
            run._element.append(elem)


def paragrafo(doc, texto, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, negrito=False,
              italico=False, tamanho=CORPO, espaco=ESPACAMENTO, depois=DEPOIS, recuo=None):
    par = doc.add_paragraph()
    _formatar_paragrafo(par, alinhamento, espaco, depois)
    if recuo is not None:
        par.paragraph_format.left_indent = recuo
    run = par.add_run(limpar(texto))
    _formatar_run(run, tamanho, negrito, italico)
    return par


def paragrafo_misto(doc, partes, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, tamanho=CORPO,
                    espaco=ESPACAMENTO, depois=DEPOIS):
    """Paragrafo com varios trechos: lista de tuplos (texto, negrito, italico)."""
    par = doc.add_paragraph()
    _formatar_paragrafo(par, alinhamento, espaco, depois)
    for texto, negrito, italico in partes:
        run = par.add_run(limpar(texto))
        _formatar_run(run, tamanho, negrito, italico)
    return par


def titulo(doc, texto, nivel=1):
    par = doc.add_heading("", level=nivel)
    _formatar_paragrafo(par, WD_ALIGN_PARAGRAPH.LEFT, ESPACAMENTO, DEPOIS,
                        Pt(12) if nivel == 1 else Pt(6))
    run = par.add_run(limpar(texto))
    _formatar_run(run, Pt(13) if nivel == 1 else CORPO, nivel <= 2)
    return par


def marca(doc, texto, tamanho=CORPO):
    """Item de lista com marca."""
    par = doc.add_paragraph(style="List Bullet")
    _formatar_paragrafo(par, WD_ALIGN_PARAGRAPH.JUSTIFY, ESPACAMENTO, Pt(3))
    run = par.add_run(limpar(texto))
    _formatar_run(run, tamanho)
    return par


def _repetir_cabecalho(linha):
    """Marca a linha como cabecalho, para se repetir quando a tabela muda de pagina."""
    tr_pr = linha._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _formatar_celula(par):
    _formatar_paragrafo(par, WD_ALIGN_PARAGRAPH.LEFT, 1.0, Pt(3))
    par.paragraph_format.left_indent = Pt(0)
    par.paragraph_format.first_line_indent = Pt(0)
    return par


def _definir_borda(celula, posicao, visivel=True, tamanho=8):
    tc_pr = celula._tc.get_or_add_tcPr()
    bordas = tc_pr.find(qn("w:tcBorders"))
    if bordas is None:
        bordas = OxmlElement("w:tcBorders")
        tc_pr.append(bordas)
    elem = bordas.find(qn(f"w:{posicao}"))
    if elem is None:
        elem = OxmlElement(f"w:{posicao}")
        bordas.append(elem)
    elem.set(qn("w:val"), "single" if visivel else "none")
    elem.set(qn("w:sz"), str(tamanho) if visivel else "0")
    elem.set(qn("w:color"), "000000")


def tabela(doc, titulo_texto, cabecalho, linhas, fonte_nota=None, tamanho=Pt(11), larguras=None):
    """Tabela de tres linhas: sem grelhas verticais, titulo acima e fonte abaixo."""
    if titulo_texto:
        paragrafo(doc, titulo_texto, WD_ALIGN_PARAGRAPH.LEFT, tamanho=Pt(11),
                  espaco=1.0, depois=Pt(3))
    tab = doc.add_table(rows=1, cols=len(cabecalho))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, texto in enumerate(cabecalho):
        celula = tab.rows[0].cells[i]
        celula.text = ""
        par = _formatar_celula(celula.paragraphs[0])
        _formatar_run(par.add_run(limpar(texto)), tamanho, negrito=True)
    _repetir_cabecalho(tab.rows[0])
    for linha in linhas:
        celulas = tab.add_row().cells
        for i, texto in enumerate(linha):
            celulas[i].text = ""
            par = _formatar_celula(celulas[i].paragraphs[0])
            _formatar_run(par.add_run(limpar(str(texto))), tamanho)
    ultima = len(tab.rows) - 1
    for idx, linha in enumerate(tab.rows):
        for celula in linha.cells:
            for lado in ("left", "right", "insideV"):
                _definir_borda(celula, lado, visivel=False)
            if idx == 0:
                _definir_borda(celula, "top")
                _definir_borda(celula, "bottom")
            elif idx == ultima:
                _definir_borda(celula, "top", visivel=False)
                _definir_borda(celula, "bottom")
            else:
                _definir_borda(celula, "top", visivel=False)
                _definir_borda(celula, "bottom", visivel=False)
    if larguras:
        for linha in tab.rows:
            for i, largura in enumerate(larguras):
                linha.cells[i].width = largura
    if fonte_nota:
        paragrafo(doc, fonte_nota, WD_ALIGN_PARAGRAPH.LEFT, tamanho=Pt(10),
                  espaco=1.0, depois=Pt(12))
    return tab


def figura(doc, caminho, legenda, largura=Cm(14)):
    """Legenda acima e imagem centrada abaixo, conforme a norma da casa."""
    par_legenda = paragrafo(doc, legenda, WD_ALIGN_PARAGRAPH.LEFT, tamanho=Pt(10),
                            espaco=1.0, depois=Pt(3))
    par_legenda.paragraph_format.keep_with_next = True
    par = doc.add_paragraph()
    _formatar_paragrafo(par, WD_ALIGN_PARAGRAPH.CENTER, 1.0, Pt(6))
    par.add_run().add_picture(caminho, width=largura)
    return par


def quebra_pagina(doc):
    par = doc.add_paragraph()
    _formatar_paragrafo(par, WD_ALIGN_PARAGRAPH.LEFT, 1.0, Pt(0))
    par.add_run().add_break(WD_BREAK.PAGE)



def subtitulo(doc, texto, tamanho=CORPO):
    """Rotulo a negrito que nao e titulo do documento, e por isso nao entra no indice."""
    par = paragrafo(doc, texto, WD_ALIGN_PARAGRAPH.LEFT, negrito=True, tamanho=tamanho,
                    espaco=1.15, depois=Pt(6))
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.keep_with_next = True
    return par


def formula(doc, partes, alinhamento=WD_ALIGN_PARAGRAPH.CENTER):
    """Formula numa linha propria, com expoentes e indices reais do Word.

    partes: lista de (texto, posicao), com posicao None, "sup" ou "sub".
    """
    par = doc.add_paragraph()
    _formatar_paragrafo(par, alinhamento, 1.0, Pt(6), Pt(3))
    par.paragraph_format.keep_with_next = True
    for texto, posicao in partes:
        run = par.add_run(limpar(texto))
        _formatar_run(run)
        run.font.superscript = posicao == "sup"
        run.font.subscript = posicao == "sub"
    return par


def _campo(par, instrucao, marcador):
    """Campo complexo do Word (inicio, instrucao, separador, resultado, fim)."""
    def fld(tipo):
        elem = OxmlElement("w:fldChar")
        elem.set(qn("w:fldCharType"), tipo)
        return elem

    run = par.add_run()
    _formatar_run(run)
    run._element.append(fld("begin"))
    texto = OxmlElement("w:instrText")
    texto.set(qn("xml:space"), "preserve")
    texto.text = instrucao
    run._element.append(texto)
    run._element.append(fld("separate"))
    resultado = par.add_run(marcador)
    _formatar_run(resultado)
    fim = par.add_run()
    fim._element.append(fld("end"))


def indice_automatico(doc, niveis="1-3"):
    """Indice automatico do Word, com os titulos dos niveis indicados e hiperligacoes.

    O python-docx nao calcula paginas: o campo e preenchido quando o Word o actualiza
    (ver actualizar_word.py). Ate la mostra apenas o marcador.
    """
    par = doc.add_paragraph()
    _formatar_paragrafo(par, WD_ALIGN_PARAGRAPH.LEFT, 1.15, Pt(0))
    _campo(par, rf' TOC \o "{niveis}" \h \z \u ', "Índice por actualizar no Word.")
    return par


def numeracao(seccao, formato, inicio=None):
    """Define o formato do numero de pagina da seccao (decimal, lowerRoman) e o recomeco."""
    sect_pr = seccao._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        # a ordem dos filhos de sectPr e fixa no esquema OOXML: pgNumType vem antes de cols
        sect_pr.insert_element_before(pg, *SUCESSORES_PGNUMTYPE)
    pg.set(qn("w:fmt"), formato)
    if inicio is not None:
        pg.set(qn("w:start"), str(inicio))
    return seccao


def nova_seccao(doc, formato="decimal", inicio=1):
    """Quebra de seccao com nova pagina e numeracao propria (por omissao, arabe desde 1)."""
    seccao = doc.add_section(WD_SECTION.NEW_PAGE)
    pg = seccao._sectPr.find(qn("w:pgNumType"))
    if pg is not None:
        seccao._sectPr.remove(pg)
    return numeracao(seccao, formato, inicio)


def referencia(doc, numero, texto):
    """Entrada Vancouver: primeira linha sem avanco, seguintes a 0,75 cm, espaco simples."""
    par = paragrafo(doc, f"{numero}. {texto}", WD_ALIGN_PARAGRAPH.JUSTIFY, espaco=1.0,
                    depois=Pt(6))
    par.paragraph_format.left_indent = AVANCO_REFERENCIAS
    par.paragraph_format.first_line_indent = -AVANCO_REFERENCIAS
    return par
