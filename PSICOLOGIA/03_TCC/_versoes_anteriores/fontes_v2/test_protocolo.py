"""Testes das funcoes com logica do protocolo: numeracao Vancouver, limpeza de texto
da casa e calculo do tamanho da amostra.

Uso: python -m pytest test_protocolo.py -q
"""
import math
import sys
from pathlib import Path

import pytest
from docx.shared import Cm

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

import house_docx as h
from _protocolo_amostra import DEFF, NAO_RESPOSTA, POPULACOES, PRECISAO, PREVALENCIA, Z, \
    amostra_base, amostra_efectiva, amostra_para, cenarios_amostra, f2_minimo_detectavel, \
    or_minimo_detectavel, parametros_maximos_logistica
from _protocolo_referencias import FONTES, Numerador


# Numeracao Vancouver

def test_atribui_numeros_pela_ordem_de_aparecimento():
    numerador = Numerador({"a": "Ref A.", "b": "Ref B.", "c": "Ref C."})
    resultado = numerador.resolver("Primeiro [[b]], depois [[a]], de novo [[b]].")
    assert resultado == "Primeiro (1), depois (2), de novo (1)."


def test_reutiliza_o_mesmo_numero_para_a_mesma_fonte():
    numerador = Numerador({"a": "Ref A."})
    numerador.resolver("Texto [[a]].")
    assert numerador.resolver("Outro texto [[a]].") == "Outro texto (1)."
    assert len(numerador.ordem) == 1


def test_comprime_tres_ou_mais_numeros_seguidos_num_intervalo():
    numerador = Numerador({"a": "A.", "b": "B.", "c": "C."})
    assert numerador.resolver("[[a,b,c]]") == "(1-3)"


def test_nao_comprime_apenas_dois_numeros_seguidos():
    numerador = Numerador({"a": "A.", "b": "B."})
    assert numerador.resolver("[[a,b]]") == "(1,2)"


def test_lista_final_segue_a_ordem_de_citacao():
    numerador = Numerador({"a": "Ref A.", "b": "Ref B."})
    numerador.resolver("[[b]] e [[a]]")
    assert numerador.lista_final() == [(1, "Ref B."), (2, "Ref A.")]


def test_chave_desconhecida_falha_de_imediato():
    numerador = Numerador({"a": "Ref A."})
    with pytest.raises(KeyError):
        numerador.resolver("[[inexistente]]")


def test_nao_citadas_identifica_fontes_orfas():
    numerador = Numerador({"a": "A.", "b": "B."})
    numerador.resolver("[[a]]")
    assert numerador.nao_citadas() == ["b"]


def test_todas_as_referencias_reais_tem_identificador_verificavel():
    for chave, texto in FONTES.items():
        assert "doi:" in texto or "Disponível em:" in texto or "PMID:" in texto, chave


def test_nenhuma_referencia_contem_travessao():
    for chave, texto in FONTES.items():
        assert "—" not in texto and "–" not in texto, chave


# Limpeza de texto no estilo da casa

def test_remove_travessoes_e_tracos():
    assert "—" not in h.limpar("a — b")
    assert h.limpar("2019–2021") == "2019-2021"


def test_repoe_ortografia_anterior_a_1990():
    assert h.limpar("o objetivo e o fator") == "o objectivo e o factor"
    assert h.limpar("uma seleção correta") == "uma selecção correcta"


def test_nao_altera_texto_ja_conforme():
    original = "O objectivo do estudo é medir factores psicossociais."
    assert h.limpar(original) == original


def test_substitui_setas_e_aproximadamente():
    assert "→" not in h.limpar("a → b")
    assert "≈" not in h.limpar("≈ 50")


# Calculo do tamanho da amostra

def test_amostra_base_corresponde_a_formula_da_proporcao_unica():
    n0, com_deff = amostra_base()
    esperado = math.ceil((Z ** 2) * PREVALENCIA * (1 - PREVALENCIA) / (PRECISAO ** 2))
    assert n0 == esperado == 383
    assert com_deff == math.ceil(esperado * DEFF)


def test_correccao_para_populacao_finita_reduz_a_amostra():
    _n0, com_deff = amostra_base()
    corrigida, _final = amostra_para(800)
    assert corrigida < com_deff


def test_amostra_bate_com_valores_de_referencia_calculados_a_mao():
    """Valores fixos, independentes da implementacao: um erro na formula seria apanhado
    mesmo que o codigo continuasse internamente consistente."""
    referencia = {800: (335, 373), 1200: (389, 433), 1600: (424, 472),
                  2000: (447, 497), 2500: (468, 520)}
    for populacao, esperado in referencia.items():
        assert amostra_para(populacao) == esperado, populacao


def test_amostra_base_tem_os_valores_de_referencia():
    assert amostra_base() == (383, 575)


def test_amostra_cresce_com_a_populacao_e_tende_para_o_limite():
    valores = [amostra_para(n)[0] for n in (800, 1200, 1600, 2000, 2500)]
    assert valores == sorted(valores)
    _n0, com_deff = amostra_base()
    assert valores[-1] < com_deff


def test_acrescimo_por_nao_resposta_e_de_dez_por_cento():
    corrigida, final = amostra_para(1600)
    assert final == math.ceil(corrigida / (1 - NAO_RESPOSTA))


def test_tabela_de_cenarios_tem_uma_linha_por_populacao_e_milhares_com_ponto():
    linhas = cenarios_amostra()
    assert len(linhas) == 5
    assert linhas[1][0] == "1.200"
    assert all(len(linha) == 3 for linha in linhas)


# Poder para o objectivo analitico. Valores de referencia obtidos por um calculo
# independente com statsmodels (NormalIndPower) e scipy (F nao central).

def test_amostra_efectiva_divide_os_respondentes_pelo_efeito_de_desenho():
    respondentes, efectiva = amostra_efectiva(1600)
    assert respondentes == 424
    assert efectiva == pytest.approx(424 / 1.5)


def test_f2_minimo_detectavel_bate_com_o_calculo_de_referencia():
    assert f2_minimo_detectavel(424 / 1.5, 15) == pytest.approx(0.0280, abs=0.0005)


def test_or_minimo_detectavel_bate_com_o_calculo_de_referencia():
    assert or_minimo_detectavel(424 / 1.5, 0.30) == pytest.approx(2.27, abs=0.02)
    assert or_minimo_detectavel(424 / 1.5, 0.20) == pytest.approx(2.50, abs=0.02)


def test_or_minimo_diminui_quando_a_exposicao_e_mais_frequente():
    valores = [or_minimo_detectavel(283, q) for q in (0.2, 0.3, 0.5)]
    assert valores == sorted(valores, reverse=True)


def test_mais_participantes_detectam_efeitos_menores():
    assert f2_minimo_detectavel(500, 15) < f2_minimo_detectavel(283, 15)
    assert or_minimo_detectavel(500, 0.3) < or_minimo_detectavel(283, 0.3)


def test_regra_de_dez_eventos_por_parametro_no_modelo_logistico():
    assert parametros_maximos_logistica(424) == 10
    assert parametros_maximos_logistica(200) == 5


# Questionario (Apendice A)

def test_numeracao_do_questionario_e_continua_e_comeca_em_um():
    from _protocolo_apendices import BLOCOS, intervalos_dos_blocos
    intervalos = intervalos_dos_blocos()
    esperado = 1
    for nome, itens in BLOCOS:
        primeiro, ultimo = intervalos[nome]
        assert primeiro == esperado, nome
        assert ultimo - primeiro + 1 == len(itens), nome
        esperado = ultimo + 1


def test_itens_nao_trazem_numeracao_escrita_a_mao():
    """A numeracao e gerada; um numero no texto do item ficaria duplicado."""
    import re
    from _protocolo_apendices import BLOCOS
    for nome, itens in BLOCOS:
        for item in itens:
            texto = item if isinstance(item, str) else item[0]
            assert not re.match(r"^\d+\.?\s", texto), f"{nome}: {texto[:40]}"


def test_questionario_regista_a_turma_para_a_analise_por_conglomerados():
    from _protocolo_apendices import SOCIODEMOGRAFICO
    assert any("turma" in pergunta.lower() for pergunta, _opcoes in SOCIODEMOGRAFICO)


def test_questionario_mede_o_indicador_usado_no_calculo_da_amostra():
    """A prevalencia de 53,7% e do indicador GSHS de violencia interpessoal: luta fisica ou
    agressao fisica sofrida nos ultimos 12 meses. As duas perguntas tem de existir."""
    from _protocolo_apendices import ESCOLA_PARES
    textos = " ".join(pergunta.lower() for pergunta, _opcoes in ESCOLA_PARES)
    assert "lutas físicas" in textos
    assert "agredido fisicamente" in textos


def test_bpaq_reduzido_tem_doze_itens_em_quatro_dimensoes():
    from _protocolo_apendices import BPAQ_ITENS
    assert len(BPAQ_ITENS) == 12
    assert {dimensao for _t, dimensao in BPAQ_ITENS} == {
        "Agressão física", "Agressão verbal", "Ira", "Hostilidade"}


# Cronograma: nenhuma actividade com participantes antes da aprovacao etica

def test_recolha_e_pre_teste_so_comecam_depois_da_aprovacao_etica():
    from _protocolo_metodologia import CRONOGRAMA

    def colunas(prefixo):
        linha = next(l for l in CRONOGRAMA if l[0].startswith(prefixo))
        return [i for i, marca in enumerate(linha[1:]) if marca]

    fim_etica = max(colunas("Submissão e aprovação"))
    assert min(colunas("Recolha de dados")) >= fim_etica
    assert min(colunas("Pré-teste")) >= fim_etica
    assert min(colunas("Análise estatística")) >= min(colunas("Recolha de dados"))


# Orcamento

def _valor(texto):
    return int(texto.replace(".", "")) if texto else 0


def test_cada_rubrica_do_orcamento_multiplica_certo():
    from _protocolo_metodologia import ORCAMENTO
    for descricao, quantidade, unitario, total in ORCAMENTO:
        digitos = "".join(c for c in quantidade if c.isdigit() or c == ".")
        if not (digitos and unitario):
            continue
        assert _valor(digitos) * _valor(unitario) == _valor(total), descricao


def test_total_do_orcamento_fecha_com_a_soma_das_rubricas():
    from _protocolo_metodologia import ORCAMENTO
    rubricas = [linha for linha in ORCAMENTO if not linha[0].startswith(("Imprevistos", "Total"))]
    subtotal = sum(_valor(linha[3]) for linha in rubricas)
    imprevistos = _valor(next(l for l in ORCAMENTO if l[0].startswith("Imprevistos"))[3])
    total = _valor(next(l for l in ORCAMENTO if l[0].startswith("Total"))[3])
    assert imprevistos == round(subtotal * 0.10)
    assert total == subtotal + imprevistos


def test_orcamento_cobre_a_amostra_prevista():
    """Ha questionarios impressos suficientes para a maior amostra dos cenarios."""
    from _protocolo_metodologia import ORCAMENTO
    impressos = _valor("".join(c for c in ORCAMENTO[0][1] if c.isdigit() or c == "."))
    maior_amostra = amostra_para(max(POPULACOES))[1]
    assert impressos >= maior_amostra * 2


# Validador: cada verificacao tem de rejeitar um documento propositadamente partido.
# Os documentos de teste usam python-docx directamente, e nao house_docx, porque este
# corrigiria os defeitos que queremos que o validador apanhe.

import validar_protocolo as v


def _documento(paragrafos):
    """Documento minimo: lista de (texto, estilo)."""
    from docx import Document
    doc = Document()
    for item in paragrafos:
        texto, estilo = item if isinstance(item, tuple) else (item, None)
        doc.add_paragraph(texto, style=estilo)
    return doc


REFERENCIAS_OK = [("REFERÊNCIAS BIBLIOGRÁFICAS", "Heading 1"),
                  "1. Autor A. Um título. Revista. 2020;1(1):1-9. doi:10.1000/a.",
                  "2. Autor B. Outro título. Revista. 2021;2(2):10-20. doi:10.1000/b."]


def test_verificar_simbolos_rejeita_travessao():
    erros = []
    v.verificar_simbolos(_documento(["Um texto com travessão — aqui."]), erros)
    assert erros


def test_verificar_ortografia_apanha_as_grafias_que_o_construtor_corrige():
    """Cada padrao corrigido pelo house_docx tem de ser detectado pelo validador."""
    casos = ["o objetivo", "um fator", "a deteção", "a reação", "a interação", "a correção",
             "a seleção", "a direção", "a infeção", "a exceção", "a adoção", "a proteção",
             "o contato", "o aspeto", "afeta o grupo", "perspetiva", "espetáculo", "correto",
             "o setor", "a atividade", "os fatores", "protecao"]
    for caso in casos:
        erros = []
        v.verificar_ortografia(_documento([f"Frase com {caso} no meio."]), erros)
        assert erros, f"grafia nao detectada: {caso}"


def test_verificar_ortografia_aceita_texto_conforme():
    erros = []
    v.verificar_ortografia(_documento(["O objectivo e os factores de protecção."]), erros)
    assert not erros


def test_verificar_vancouver_rejeita_citacao_acima_do_maior_numero_listado():
    """Regressao: o filtro por maximo escondia o caso mais obvio de citacao invalida."""
    doc = _documento(["Uma afirmação com citação inexistente (60)."] + REFERENCIAS_OK)
    erros, avisos = [], []
    v.verificar_vancouver(doc, erros, avisos)
    assert any("sem entrada na lista" in e for e in erros)


def test_verificar_vancouver_rejeita_saltos_na_lista():
    doc = _documento(["Texto com citação (1).", ("REFERÊNCIAS BIBLIOGRÁFICAS", "Heading 1"),
                      "1. Autor A. Título. Revista. 2020;1(1):1-9. doi:10.1000/a.",
                      "3. Autor C. Título. Revista. 2022;3(3):30-9. doi:10.1000/c."])
    erros, avisos = [], []
    v.verificar_vancouver(doc, erros, avisos)
    assert any("saltos" in e for e in erros)


def test_verificar_vancouver_aceita_lista_coerente():
    doc = _documento(["Texto com citações (1,2)."] + REFERENCIAS_OK)
    erros, avisos = [], []
    v.verificar_vancouver(doc, erros, avisos)
    assert not erros


def test_verificar_vancouver_ignora_numeros_dentro_das_referencias():
    """O volume e as paginas de uma referencia nao sao citacoes."""
    doc = _documento(["Texto com citação (1).", ("REFERÊNCIAS BIBLIOGRÁFICAS", "Heading 1"),
                      "1. von Elm E. STROBE. Lancet. 2007;370(9596):1453-7. doi:10.1000/x."])
    erros, avisos = [], []
    v.verificar_vancouver(doc, erros, avisos)
    assert not erros


def _doc_objectivos_resultados(n_objectivos, n_resultados, com_frase_introdutoria=True):
    paragrafos = [("3.2. Objectivos específicos", "Heading 2")]
    paragrafos += [(f"Objectivo específico número {i}, redigido com verbo no infinitivo e "
                    f"formulado de modo mensurável.", "List Bullet")
                   for i in range(n_objectivos)]
    paragrafos += [("4. HIPÓTESES DE INVESTIGAÇÃO", "Heading 1"),
                   ("8. RESULTADOS ESPERADOS", "Heading 1")]
    if com_frase_introdutoria:
        paragrafos.append("Em coerência com os objectivos específicos, esperam-se os resultados.")
    paragrafos += [(f"Resultado esperado número {i}, correspondente ao objectivo específico "
                    f"com o mesmo número de ordem.", "List Bullet")
                   for i in range(n_resultados)]
    paragrafos.append(("9. DIVULGAÇÃO DOS RESULTADOS", "Heading 1"))
    return _documento(paragrafos)


def test_verificar_coerencia_rejeita_menos_resultados_que_objectivos():
    erros = []
    v.verificar_coerencia(_doc_objectivos_resultados(5, 4), erros)
    assert any("regra de ouro" in e for e in erros)


def test_verificar_coerencia_nao_conta_a_frase_introdutoria_como_resultado():
    """Regressao: a frase de abertura da seccao 8 inflacionava a contagem."""
    erros = []
    v.verificar_coerencia(_doc_objectivos_resultados(5, 4, com_frase_introdutoria=True), erros)
    assert erros, "a frase introdutoria esta a mascarar um resultado em falta"


def test_verificar_coerencia_aceita_correspondencia_completa():
    erros = []
    v.verificar_coerencia(_doc_objectivos_resultados(5, 5), erros)
    assert not erros


def test_verificar_abreviaturas_rejeita_sigla_sem_uso_no_corpo():
    doc = _documento([("LISTA DE ABREVIATURAS E SIGLAS", "Heading 1"),
                      "OMS = Organização Mundial da Saúde;",
                      "XPTO = sigla que ninguém usa;",
                      ("ÍNDICE", "Heading 1"),
                      "A OMS recomenda esta abordagem no corpo do texto."])
    erros = []
    v.verificar_abreviaturas(doc, erros)
    assert any("XPTO" in e for e in erros)


def _doc_siglas(*corpo):
    return _documento([("LISTA DE ABREVIATURAS E SIGLAS", "Heading 1"),
                       "OMS = Organização Mundial da Saúde;", "BPAQ = Buss-Perry Aggression "
                       "Questionnaire;", ("ÍNDICE", None), ("1. INTRODUÇÃO", "Heading 1")]
                      + list(corpo))


def test_verificar_abreviaturas_rejeita_sigla_usada_antes_da_forma_extensa():
    doc = _doc_siglas("A OMS recomenda.", "A Organização Mundial da Saúde (OMS) recomenda.",
                      "O Buss-Perry Aggression Questionnaire (BPAQ) mede a agressividade.")
    erros = []
    v.verificar_abreviaturas(doc, erros)
    assert any("primeira ocorrencia" in e and "OMS" in e for e in erros)


def test_verificar_abreviaturas_nao_confunde_sigla_com_parte_de_outra():
    """Regressao: BPAQ era dado como usado so por aparecer dentro de BPAQ-SF."""
    doc = _doc_siglas("A Organização Mundial da Saúde (OMS) recomenda.", "A forma BPAQ-SF.")
    erros = []
    v.verificar_abreviaturas(doc, erros)
    assert any("BPAQ" in e and "ausentes" in e for e in erros)


def test_verificar_abreviaturas_rejeita_sigla_perdida_dentro_de_parenteses_alheios():
    """A sigla tem de fechar o parentese da forma extensa, nao apenas estar dentro de um."""
    doc = _doc_siglas("Várias entidades (entre elas a OMS e outras agências da ONU) recomendam.",
                      "O Buss-Perry Aggression Questionnaire (BPAQ) mede a agressividade.")
    erros = []
    v.verificar_abreviaturas(doc, erros)
    assert any("primeira ocorrencia" in e and "OMS" in e for e in erros)


def test_verificar_abreviaturas_aceita_sigla_seguida_de_qualificador():
    """IC95% e a forma usada para o intervalo de confianca a 95 por cento."""
    doc = _documento([("LISTA DE ABREVIATURAS E SIGLAS", "Heading 1"),
                      "IC = intervalo de confiança;", ("ÍNDICE", None),
                      ("1. INTRODUÇÃO", "Heading 1"),
                      "Intervalos de confiança a 95 por cento (IC95%)."])
    erros = []
    v.verificar_abreviaturas(doc, erros)
    assert not erros, erros


def test_verificar_abreviaturas_aceita_siglas_definidas_na_primeira_ocorrencia():
    doc = _doc_siglas("A Organização Mundial da Saúde (OMS) recomenda.",
                      "O questionário (Buss-Perry Aggression Questionnaire, BPAQ) e a OMS.")
    erros = []
    v.verificar_abreviaturas(doc, erros)
    assert not erros, erros


def test_verificar_estrutura_rejeita_seccao_em_falta():
    erros = []
    v.verificar_estrutura(_documento([("1. INTRODUÇÃO", "Heading 1")]), erros)
    assert any("7. METODOLOGIA" in e for e in erros)


def test_verificar_rodape_rejeita_documento_sem_numero_de_pagina():
    from docx import Document
    erros = []
    v.verificar_rodape(Document(), erros)
    assert erros


def test_verificar_pagina_rejeita_margens_erradas():
    from docx import Document
    from docx.shared import Cm
    doc = Document()
    doc.sections[0].top_margin = Cm(1)
    erros = []
    v.verificar_pagina(doc, erros)
    assert any("topo" in e for e in erros)


def test_verificar_coerencia_ignora_as_entradas_do_indice():
    """Regressao: com indice automatico, o texto '3.2. Objectivos especificos' aparece
    primeiro no indice, fora de qualquer lista, e escondia os objectivos reais."""
    doc = _doc_objectivos_resultados(5, 5)
    primeiro = doc.paragraphs[0]
    for texto in ("3.2. Objectivos específicos\t8", "4. HIPÓTESES DE INVESTIGAÇÃO\t9"):
        primeiro.insert_paragraph_before(texto)
    erros = []
    v.verificar_coerencia(doc, erros)
    assert not erros


RESUMO_VALIDO = " ".join(["palavra"] * 270) + "."


def _doc_resumo(resumo=RESUMO_VALIDO, chaves="adolescente, agressividade, violência.",
                abstract="An abstract written as a single paragraph.",
                keywords="adolescent, aggression, violence."):
    return _documento([("RESUMO", "Heading 1"), resumo, f"Palavras-chave: {chaves}",
                       ("ABSTRACT", "Heading 1"), abstract, f"Keywords: {keywords}"])


def test_verificar_resumo_aceita_resumo_conforme():
    erros = []
    v.verificar_resumo(_doc_resumo(), erros)
    assert not erros, erros


def test_verificar_resumo_rejeita_rotulos_em_serie():
    erros = []
    v.verificar_resumo(_doc_resumo("Introdução: texto. Objectivo: " + RESUMO_VALIDO), erros)
    assert any("rotulos" in e for e in erros)


def test_verificar_resumo_rejeita_extensao_fora_de_250_a_300_palavras():
    erros = []
    v.verificar_resumo(_doc_resumo(" ".join(["palavra"] * 320)), erros)
    assert any("palavras" in e for e in erros)


def test_verificar_resumo_rejeita_siglas_e_citacoes():
    for intruso in ("A OMS recomenda.", "Como se viu (12)."):
        erros = []
        v.verificar_resumo(_doc_resumo(intruso + " " + RESUMO_VALIDO), erros)
        assert erros, intruso


def test_verificar_resumo_rejeita_mais_de_cinco_palavras_chave():
    erros = []
    v.verificar_resumo(_doc_resumo(chaves="a, b, c, d, e, f."), erros)
    assert any("palavras-chave" in e for e in erros)


def test_verificar_resumo_rejeita_palavras_chave_fora_de_ordem_alfabetica():
    erros = []
    v.verificar_resumo(_doc_resumo(chaves="violência, adolescente, agressividade."), erros)
    assert any("alfabetica" in e for e in erros)


def test_verificar_resumo_aceita_acentos_na_ordem_alfabetica():
    erros = []
    v.verificar_resumo(_doc_resumo(chaves="adolescente, factores, Moçambique, violência."),
                       erros)
    assert not erros, erros


def test_verificar_indice_rejeita_indice_manual():
    erros = []
    v.verificar_indice(_documento([("ÍNDICE", None), "1. Introdução"]), erros)
    assert any("automatico" in e for e in erros)


def test_verificar_indice_rejeita_campo_por_actualizar():
    doc = _documento([("ÍNDICE", None)])
    h.indice_automatico(doc)
    doc.add_paragraph("1. INTRODUÇÃO", style="Heading 1")
    erros = []
    v.verificar_indice(doc, erros)
    assert any("actualizar" in e for e in erros)


def test_verificar_indice_aceita_campo_actualizado():
    doc = _documento([("ÍNDICE", None)])
    h.indice_automatico(doc)
    doc.add_paragraph("1. INTRODUÇÃO\t1")
    doc.add_paragraph("1. INTRODUÇÃO", style="Heading 1")
    erros = []
    v.verificar_indice(doc, erros)
    assert not erros, erros


def test_verificar_simbolos_rejeita_expoente_unicode():
    erros = []
    v.verificar_simbolos(_documento(["n = Z² × p"]), erros)
    assert any("expoente" in e for e in erros)


def test_verificar_referencias_rejeita_entrada_sem_avanco():
    erros = []
    v.verificar_referencias(_documento(REFERENCIAS_OK), erros)
    assert any("0,75" in e for e in erros)


def test_verificar_referencias_aceita_entrada_com_avanco():
    doc = h.novo_documento()
    h.referencia(doc, 1, "Autor A. Um título. Revista. 2020;1(1):1-9. doi:10.1000/a.")
    erros = []
    v.verificar_referencias(doc, erros)
    assert not erros, erros


def test_verificar_rodape_rejeita_rodape_fora_de_1_5_cm():
    doc = h.novo_documento()
    h.numerar_paginas(doc)
    doc.sections[0].footer_distance = Cm(1.25)
    erros = []
    v.verificar_rodape(doc, erros)
    assert any("1,5" in e for e in erros)


def test_validador_aprova_o_protocolo_real():
    caminho = BASE / "03_TCC" / "PROTOCOLO_AGRESSIVIDADE_MUATALA.docx"
    if not caminho.exists():
        pytest.skip("protocolo ainda nao construido")
    assert v.validar(str(caminho)) == 0


# Documento construido

@pytest.fixture(scope="module")
def documento():
    from docx import Document
    caminho = BASE / "03_TCC" / "PROTOCOLO_AGRESSIVIDADE_MUATALA.docx"
    if not caminho.exists():
        pytest.skip("protocolo ainda nao construido")
    return Document(caminho)


def test_documento_tem_as_margens_e_o_formato_da_casa(documento):
    from docx.shared import Cm
    sec = documento.sections[0]
    assert abs(sec.page_width - Cm(21)) < 5000
    assert abs(sec.top_margin - Cm(3)) < 5000
    assert abs(sec.left_margin - Cm(2.5)) < 5000


def test_documento_nao_contem_simbolos_proibidos(documento):
    proibidos = ("—", "–", "→", "≈")
    for par in documento.paragraphs:
        assert not any(s in par.text for s in proibidos), par.text[:80]


def test_documento_tem_todas_as_seccoes_obrigatorias(documento):
    titulos = " | ".join(p.text for p in documento.paragraphs
                         if p.style.name.startswith("Heading"))
    for seccao in ("1. INTRODUÇÃO", "7. METODOLOGIA", "11. ORÇAMENTO",
                   "REFERÊNCIAS BIBLIOGRÁFICAS", "APÊNDICES"):
        assert seccao in titulos


def test_documento_numera_pre_textuais_em_romano_e_corpo_em_arabe_desde_um(documento):
    from docx.oxml.ns import qn
    primeira = documento.sections[0]._sectPr.find(qn("w:pgNumType"))
    ultima = documento.sections[-1]._sectPr.find(qn("w:pgNumType"))
    assert primeira.get(qn("w:fmt")) == "lowerRoman"
    # o Word, ao gravar, apaga fmt="decimal" por ser o formato por omissao
    assert ultima.get(qn("w:fmt")) in (None, "decimal")
    assert ultima.get(qn("w:start")) == "1"


def test_subseccoes_numeradas_de_segundo_nivel_sao_titulos_de_nivel_2(documento):
    """5.1 a 5.4 sao do mesmo nivel que 2.1 e 7.1; nivel 3 fica para 7.5.1 e 7.5.2."""
    import re
    for par in documento.paragraphs:
        if re.match(r"^\d+\.\d+\.\s", par.text) and par.style.name.startswith("Heading"):
            assert par.style.name == "Heading 2", par.text


def test_quadro_1_identifica_cada_estudo_pelo_numero_vancouver(documento):
    import re
    quadro = next(t for t in documento.tables if t.rows[0].cells[0].text == "Autor (ano)")
    for linha in quadro.rows[1:]:
        assert re.search(r"\(\d+\)$", linha.cells[0].text), linha.cells[0].text


def test_formulas_da_amostra_usam_expoentes_reais(documento):
    runs = [r for p in documento.paragraphs for r in p.runs]
    assert any(r.font.superscript and r.text == "2" for r in runs)
    assert any(r.font.subscript and r.text == "0" for r in runs)
