"""Valida o protocolo contra as regras da casa UniLurio antes da entrega.

Verifica formatacao (pagina, fonte, tamanho, cor, espacamento, alinhamento), ausencia
de travessoes e de marcas de geracao automatica, ortografia anterior ao Acordo de 1990,
integridade da numeracao Vancouver e correspondencia entre objectivos e resultados
esperados.

Uso: python validar_protocolo.py [caminho.docx]
"""
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph

import house_docx

PROIBIDOS = {"—": "travessao (em dash)", "–": "traco (en dash)",
             "−": "sinal de menos", "→": "seta", "≈": "aproximadamente",
             "…": "reticencias tipograficas"}
PROIBIDOS.update({c: "expoente ou indice Unicode, usar sobrescrito ou subscrito do Word"
                  for c in "⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉"})

# Norma da FCS: resumo em paragrafo unico de 250 a 300 palavras, sem citacoes, sem
# abreviaturas e com 3 a 5 palavras-chave por ordem alfabetica.
PALAVRAS_RESUMO = (250, 300)
PALAVRAS_CHAVE = (3, 5)
ROTULOS = re.compile(r"(?:^|[.!?]\s+)(Introdução|Objectivos?|Metodologia|Métodos|Resultados"
                     r"(?: esperados)?|Conclus(?:ão|ões)|Background|Objectives?|Methods|"
                     r"Results|Expected results|Conclusions?)\s*:", re.IGNORECASE)
SIGLA = re.compile(r"\b[A-ZÁÂÃÀÇÉÊÍÓÔÕÚ]{2,}\b")
CITACAO = re.compile(r"\((\d+(?:[,-]\d+)*)\)")

# Os padroes de deteccao sao exactamente os que o construtor corrige, importados de
# house_docx para que o detector e o corrector nao possam divergir.
POS_1990 = [padrao for padrao, _correccao in house_docx.PRE_1990]

CLICHES = ["em suma", "vale ressaltar", "é importante notar", "mergulhar", "em resumo,",
           "por fim, é", "cabe destacar", "vale a pena notar", "como modelo de linguagem"]

FONTE = "Times New Roman"


def _texto_de(doc):
    """Todos os textos do documento: paragrafos e celulas de tabelas."""
    for par in doc.paragraphs:
        if par.text.strip():
            yield ("paragrafo", par.text)
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                if celula.text.strip():
                    yield ("tabela", celula.text)


def _runs_de(doc):
    for par in doc.paragraphs:
        for run in par.runs:
            if run.text.strip():
                yield par, run
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                for par in celula.paragraphs:
                    for run in par.runs:
                        if run.text.strip():
                            yield par, run


def verificar_pagina(doc, erros):
    sec = doc.sections[0]
    esperado = {"largura": Cm(21), "altura": Cm(29.7), "topo": Cm(3), "base": Cm(3),
                "esquerda": Cm(2.5), "direita": Cm(2.5)}
    obtido = {"largura": sec.page_width, "altura": sec.page_height, "topo": sec.top_margin,
              "base": sec.bottom_margin, "esquerda": sec.left_margin, "direita": sec.right_margin}
    for chave, valor in esperado.items():
        if abs(obtido[chave] - valor) > 5000:
            erros.append(f"pagina: {chave} = {obtido[chave].cm:.2f} cm, esperado {valor.cm:.2f} cm")


def _herdado(doc, par, ler):
    """Valor efectivo de um atributo de letra: estilo do paragrafo, os seus estilos-base e,
    por fim, o Normal. O Word, ao gravar, apaga a formatacao directa igual a do estilo."""
    estilo = par.style
    while estilo is not None:
        valor = ler(estilo.font)
        if valor is not None:
            return valor
        estilo = estilo.base_style
    return ler(doc.styles["Normal"].font)


def _cor(fonte):
    return fonte.color.rgb if fonte.color is not None and fonte.color.type is not None else None


def verificar_formatacao(doc, erros, avisos):
    fontes, tamanhos, cores = Counter(), Counter(), Counter()
    for par, run in _runs_de(doc):
        nome = run.font.name or _herdado(doc, par, lambda f: f.name)
        fontes[nome] += 1
        tamanho = run.font.size or _herdado(doc, par, lambda f: f.size)
        tamanhos[tamanho.pt if tamanho else None] += 1
        cor = _cor(run.font) or _herdado(doc, par, _cor)
        cores[str(cor)] += 1
    for nome, quantos in fontes.items():
        if nome != FONTE:
            erros.append(f"fonte diferente de {FONTE}: {nome} em {quantos} trechos")
    for cor, quantos in cores.items():
        if cor not in ("000000", "None"):
            erros.append(f"cor diferente de preto: {cor} em {quantos} trechos")
        if cor == "None":
            avisos.append(f"{quantos} trechos sem cor explicita, herdam o estilo Normal")
    corpo = tamanhos.get(12.0, 0)
    if corpo == 0:
        erros.append("nenhum trecho a 12 pt, o corpo de texto deve estar a 12 pt")
    fora = {t: q for t, q in tamanhos.items() if t not in (9.5, 10.0, 11.0, 12.0, 13.0, 14.0)}
    if fora:
        erros.append(f"tamanhos de letra inesperados: {fora}")


def verificar_espacamento(doc, erros):
    problemas = 0
    for par in doc.paragraphs:
        if not par.text.strip():
            continue
        espaco = par.paragraph_format.line_spacing
        if espaco is not None and espaco not in (1.0, 1.15, 1.5):
            problemas += 1
    if problemas:
        erros.append(f"{problemas} paragrafos com espacamento fora de 1,0, 1,15 ou 1,5")


def verificar_simbolos(doc, erros):
    for origem, texto in _texto_de(doc):
        for simbolo, nome in PROIBIDOS.items():
            if simbolo in texto:
                erros.append(f"{nome} encontrado em {origem}: {texto[:90]}")


def verificar_ortografia(doc, erros):
    for origem, texto in _texto_de(doc):
        minusculas = texto.lower()
        for padrao in POS_1990:
            achado = re.search(padrao, minusculas)
            if achado:
                erros.append(f"grafia pos-1990 '{achado.group(0)}' em {origem}: {texto[:90]}")


def verificar_cliches(doc, avisos):
    for origem, texto in _texto_de(doc):
        minusculas = texto.lower()
        for cliche in CLICHES:
            if cliche in minusculas:
                avisos.append(f"cliche '{cliche}' em {origem}: {texto[:90]}")


def _e_entrada_de_referencia(texto):
    """Distingue uma entrada da lista final de um paragrafo do corpo do texto."""
    return bool(re.match(r"^\d+\.\s", texto.strip())) and (
        "doi:" in texto or "Disponível em:" in texto or "PMID:" in texto)


def verificar_vancouver(doc, erros, avisos):
    citados = set()
    for _origem, texto in _texto_de(doc):
        if _e_entrada_de_referencia(texto):
            continue
        for bloco in CITACAO.findall(texto):
            for parte in bloco.split(","):
                if "-" in parte:
                    principio, fim = parte.split("-")
                    citados.update(range(int(principio), int(fim) + 1))
                else:
                    citados.add(int(parte))
    listados = set()
    for par in doc.paragraphs:
        if _e_entrada_de_referencia(par.text):
            listados.add(int(re.match(r"^(\d+)\.", par.text.strip()).group(1)))
    if not listados:
        erros.append("nao foi encontrada a lista de referencias")
        return
    maximo = max(listados)
    if listados != set(range(1, maximo + 1)):
        erros.append(f"lista de referencias com saltos: faltam "
                     f"{sorted(set(range(1, maximo + 1)) - listados)}")
    sem_lista = citados - listados
    if sem_lista:
        erros.append(f"numeros citados sem entrada na lista de referencias: "
                     f"{sorted(sem_lista)}")
    nao_citadas = listados - citados
    if nao_citadas:
        avisos.append(f"referencias listadas sem citacao detectada no texto: {sorted(nao_citadas)}")
    print(f"  referencias listadas: {maximo}; numeros distintos citados no texto: {len(citados)}")


def _texto_em_ordem(doc):
    """Paragrafos e celulas de tabela pela ordem em que aparecem no documento."""
    for elemento in doc.element.body.iterchildren():
        if elemento.tag == qn("w:p"):
            par = Paragraph(elemento, doc)
            yield par.text, par.style.name
        elif elemento.tag == qn("w:tbl"):
            for linha in Table(elemento, doc).rows:
                for celula in linha.cells:
                    yield celula.text, "Tabela"


def _corpo_do_texto(doc):
    """Texto a partir da Introducao (ou do indice, se nao houver Introducao)."""
    itens = list(_texto_em_ordem(doc))
    inicio = next((i for i, (t, s) in enumerate(itens)
                   if s.startswith("Heading") and t.startswith("1. INTRODUÇÃO")), None)
    if inicio is None:
        inicio = next((i for i, (t, _s) in enumerate(itens) if t.strip() == "ÍNDICE"), 0)
    return "\n".join(t for t, _s in itens[inicio + 1:])


def _fecha_parenteses_da_forma_extensa(texto, inicio, fim):
    """A sigla esta dentro de um parentese aberto ha pouco e fecha-o logo a seguir, como em
    'Forma extensa (SIGLA)', 'Nome original, SIGLA)' ou '(IC95%)'."""
    abre, fecha = texto.rfind("(", 0, inicio), texto.rfind(")", 0, inicio)
    if abre <= fecha or inicio - abre > 120:
        return False
    seguinte = re.match(r"[^()\s]{0,6}\)", texto[fim:])
    return seguinte is not None


def verificar_abreviaturas(doc, erros):
    """Cada sigla listada aparece no corpo e, na primeira vez, entre parenteses a seguir a
    forma extensa, como exige a norma da FCS."""
    textos = [par.text.strip() for par in doc.paragraphs]
    try:
        inicio = next(i for i, t in enumerate(textos) if t.startswith("LISTA DE ABREVIATURAS"))
        fim = next(i for i, t in enumerate(textos) if i > inicio and t.startswith("ÍNDICE"))
    except StopIteration:
        erros.append("nao foi encontrada a lista de abreviaturas")
        return
    siglas = [t.split("=")[0].strip() for t in textos[inicio + 1:fim] if "=" in t]
    corpo = _corpo_do_texto(doc)
    ausentes, sem_forma_extensa = [], []
    for sigla in siglas:
        achado = re.search(r"(?<![A-Za-zÀ-ü-])" + re.escape(sigla) + r"(?![A-Za-zÀ-ü-])", corpo)
        if achado is None:
            ausentes.append(sigla)
        elif not _fecha_parenteses_da_forma_extensa(corpo, achado.start(), achado.end()):
            sem_forma_extensa.append(sigla)
    if ausentes:
        erros.append(f"siglas listadas mas ausentes do corpo do texto: {', '.join(ausentes)}")
    if sem_forma_extensa:
        erros.append("siglas cuja primeira ocorrencia nao vem entre parenteses a seguir a forma "
                     f"extensa: {', '.join(sem_forma_extensa)}")
    print(f"  siglas na lista de abreviaturas: {len(siglas)}")


def verificar_coerencia(doc, erros):
    """A cada objectivo especifico deve corresponder, no minimo, um resultado esperado.

    Conta apenas itens de lista, para que frases introdutorias das seccoes nao inflacionem
    a contagem e escondam um desalinhamento real. Os limites das seccoes sao procurados
    so entre os titulos, porque o indice automatico repete o mesmo texto mais acima.
    """
    pares = [(par.text.strip(), par.style.name) for par in doc.paragraphs]

    def marcas_entre(inicio, fim):
        try:
            i = next(k for k, (t, s) in enumerate(pares)
                     if s.startswith("Heading") and t.startswith(inicio))
            j = next(k for k, (t, s) in enumerate(pares)
                     if k > i and s.startswith("Heading") and t.startswith(fim))
        except StopIteration:
            return []
        return [texto for texto, estilo in pares[i + 1:j]
                if estilo.startswith("List") and len(texto) > 40]

    objectivos = marcas_entre("3.2. Objectivos específicos", "4. HIPÓTESES")
    resultados = marcas_entre("8. RESULTADOS ESPERADOS", "9. DIVULGAÇÃO")
    if not objectivos:
        erros.append("nao foram encontrados objectivos especificos")
    if not resultados:
        erros.append("nao foram encontrados resultados esperados")
    if resultados and len(resultados) < len(objectivos):
        erros.append(f"{len(objectivos)} objectivos especificos para apenas {len(resultados)} "
                     f"resultados esperados, a regra de ouro exige correspondencia")
    print(f"  objectivos especificos: {len(objectivos)}; resultados esperados: {len(resultados)}")


def verificar_estrutura(doc, erros):
    obrigatorias = ["1. INTRODUÇÃO", "2. IDENTIFICAÇÃO", "3. OBJECTIVOS", "4. HIPÓTESES",
                    "5. JUSTIFICATIVA", "6. REVISÃO DA LITERATURA", "7. METODOLOGIA",
                    "8. RESULTADOS ESPERADOS", "9. DIVULGAÇÃO", "10. CRONOGRAMA", "11. ORÇAMENTO",
                    "12. RECURSOS HUMANOS", "REFERÊNCIAS BIBLIOGRÁFICAS", "APÊNDICES"]
    titulos = [par.text.strip() for par in doc.paragraphs if par.style.name.startswith("Heading")]
    for seccao in obrigatorias:
        if not any(t.startswith(seccao) for t in titulos):
            erros.append(f"seccao obrigatoria em falta: {seccao}")


def verificar_rodape(doc, erros):
    for sec in doc.sections:
        xml = sec.footer.paragraphs[0]._element.xml if sec.footer.paragraphs else ""
        if "PAGE" not in xml:
            erros.append("o rodape nao contem o campo de numero de pagina")
        elif "w:jc w:val=\"right\"" not in sec.footer.paragraphs[0]._element.xml:
            erros.append("o numero de pagina nao esta alinhado a direita")
        for nome, distancia in (("cabecalho", sec.header_distance),
                                ("rodape", sec.footer_distance)):
            if distancia is None or abs(distancia - house_docx.DISTANCIA_CABECALHO_RODAPE) > 5000:
                erros.append(f"o {nome} deve ficar a 1,5 cm da margem da folha")


def _ordem_alfabetica(termo):
    sem_acentos = unicodedata.normalize("NFKD", termo)
    return "".join(c for c in sem_acentos if not unicodedata.combining(c)).casefold()


def _bloco_de_resumo(doc, titulo, rotulo):
    """Paragrafos do corpo do resumo e a linha de palavras-chave que o fecha."""
    pares = doc.paragraphs
    inicio = next((i for i, p in enumerate(pares)
                   if p.style.name.startswith("Heading") and p.text.strip() == titulo), None)
    if inicio is None:
        return None, None
    corpo = []
    for par in pares[inicio + 1:]:
        texto = par.text.strip()
        if texto.startswith(rotulo):
            return corpo, texto
        if par.style.name.startswith("Heading"):
            break
        if texto:
            corpo.append(texto)
    return corpo, None


def _verificar_palavras_chave(titulo, linha, erros):
    if linha is None:
        erros.append(f"{titulo}: faltam as palavras-chave")
        return
    termos = [t.strip(" .;") for t in linha.split(":", 1)[1].split(",") if t.strip(" .;")]
    minimo, maximo = PALAVRAS_CHAVE
    if not minimo <= len(termos) <= maximo:
        erros.append(f"{titulo}: {len(termos)} palavras-chave, a norma pede {minimo} a {maximo}")
    if termos != sorted(termos, key=_ordem_alfabetica):
        erros.append(f"{titulo}: palavras-chave fora de ordem alfabetica")


def verificar_resumo(doc, erros):
    """Resumo e abstract: paragrafo unico, prosa corrida, sem siglas nem citacoes."""
    for titulo, rotulo, extensao in (("RESUMO", "Palavras-chave:", PALAVRAS_RESUMO),
                                     ("ABSTRACT", "Keywords:", None)):
        corpo, linha_chave = _bloco_de_resumo(doc, titulo, rotulo)
        if not corpo:
            erros.append(f"{titulo}: nao foi encontrado")
            continue
        if len(corpo) != 1:
            erros.append(f"{titulo}: deve ser um paragrafo unico, tem {len(corpo)}")
        texto = " ".join(corpo)
        if extensao:
            palavras = len(texto.split())
            if not extensao[0] <= palavras <= extensao[1]:
                erros.append(f"{titulo}: {palavras} palavras, a norma pede "
                             f"{extensao[0]} a {extensao[1]}")
        if ROTULOS.search(texto):
            erros.append(f"{titulo}: rotulos em serie (Introducao:, Objectivo:...), "
                         "escrever em prosa corrida")
        if CITACAO.search(texto):
            erros.append(f"{titulo}: contem citacoes, a norma proibe-as no resumo")
        siglas = sorted(set(SIGLA.findall(texto)))
        if siglas:
            erros.append(f"{titulo}: contem siglas ({', '.join(siglas)}), a norma proibe-as")
        _verificar_palavras_chave(titulo, linha_chave, erros)


def verificar_indice(doc, erros):
    """O indice tem de ser o campo automatico do Word, ja actualizado com as paginas."""
    corpo = doc.element.body
    instrucoes = [n.text or "" for n in corpo.iter(qn("w:instrText"))]
    instrucoes += [n.get(qn("w:instr")) or "" for n in corpo.iter(qn("w:fldSimple"))]
    if not any("TOC" in i for i in instrucoes):
        erros.append("o indice nao e automatico: inserir o campo de indice do Word")
        return
    textos = [par.text for par in doc.paragraphs]
    if not any(re.search(r"INTRODUÇÃO\s*\t\s*\d+", t) for t in textos):
        erros.append("o indice automatico esta por actualizar: correr actualizar_word.py")


def verificar_referencias(doc, erros):
    """Entradas com a primeira linha sem avanco, as restantes a 0,75 cm e espaco simples."""
    alvo = house_docx.AVANCO_REFERENCIAS
    sem_avanco = espacamento = 0
    for par in doc.paragraphs:
        if not _e_entrada_de_referencia(par.text):
            continue
        pf = par.paragraph_format
        if (pf.left_indent is None or pf.first_line_indent is None
                or abs(pf.left_indent - alvo) > 5000 or abs(pf.first_line_indent + alvo) > 5000):
            sem_avanco += 1
        if pf.line_spacing not in (None, 1.0):
            espacamento += 1
    if sem_avanco:
        erros.append(f"{sem_avanco} referencias sem avanco de 0,75 cm nas linhas seguintes")
    if espacamento:
        erros.append(f"{espacamento} referencias sem espaco simples dentro da entrada")


def validar(caminho):
    doc = Document(caminho)
    erros, avisos = [], []
    print(f"Validacao de {Path(caminho).name}")
    verificar_pagina(doc, erros)
    verificar_formatacao(doc, erros, avisos)
    verificar_espacamento(doc, erros)
    verificar_simbolos(doc, erros)
    verificar_ortografia(doc, erros)
    verificar_cliches(doc, avisos)
    verificar_estrutura(doc, erros)
    verificar_resumo(doc, erros)
    verificar_indice(doc, erros)
    verificar_abreviaturas(doc, erros)
    verificar_vancouver(doc, erros, avisos)
    verificar_referencias(doc, erros)
    verificar_coerencia(doc, erros)
    verificar_rodape(doc, erros)

    palavras = sum(len(t.split()) for _o, t in _texto_de(doc))
    print(f"  paragrafos: {len(doc.paragraphs)}; tabelas: {len(doc.tables)}; palavras: {palavras}")
    for aviso in avisos:
        print(f"  AVISO: {aviso}")
    if erros:
        print(f"\n{len(erros)} ERRO(S):")
        for erro in erros:
            print(f"  - {erro}")
        return 1
    print("\nSem erros: o documento cumpre as regras da casa.")
    return 0


if __name__ == "__main__":
    alvo = sys.argv[1] if len(sys.argv) > 1 else str(
        Path(__file__).resolve().parent / "03_TCC" / "PROTOCOLO_AGRESSIVIDADE_MUATALA.docx")
    sys.exit(validar(alvo))
