# -*- coding: utf-8 -*-
"""
Primitivas de composicao na norma da Faculdade de Ciencias de Saude da
Universidade Lurio (documentos 00, 01 e 12 da pasta PROMPTS_TCC_UNILURIO),
com o padrao de documentos do utilizador onde a norma e' omissa.

  Pagina A4; margens 3 cm (topo e base) e 2,5 cm (esquerda e direita);
  cabecalho e rodape a 1,5 cm; so o numero de pagina, no canto inferior
  direito; a capa nao mostra numero.
  Times New Roman 12, preto; espacamento 1,5; 6 pt depois; justificado.
  Titulos: nivel 1 a 13 pt negrito, nivel 2 a 12 pt negrito, nivel 3 sem
  negrito. Tabelas a 11 pt, com tres linhas horizontais e sem grelha
  vertical; legenda acima (10 pt) e fonte abaixo (10 pt), no formato
  "Tabela 1: titulo".
  Referencias: primeira linha sem avanco, restantes com 0,75 cm.
"""
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONTE = "Times New Roman"
CORPO = Pt(12)
ESPACAMENTO = 1.5
DEPOIS = Pt(6)
PRETO = RGBColor(0, 0, 0)
LARGURA_UTIL_CM = 16.0
TAM_TABELA = 11
TAM_LEGENDA = 10

AL = WD_ALIGN_PARAGRAPH

# --------------------------------------------------------------------------
# documento e estilos
# --------------------------------------------------------------------------


def novo_documento():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(3.0)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header_distance = sec.footer_distance = Cm(1.5)
    sec.different_first_page_header_footer = True
    _estilos(doc)
    _numero_de_pagina(sec)
    return doc


def _fonte_estilo(st, tamanho, negrito=False):
    st.font.name = FONTE
    st.font.size = tamanho
    st.font.bold = negrito
    st.font.italic = False
    st.font.color.rgb = PRETO
    rpr = st._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(att), FONTE)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme",
                "w:cstheme"):
        if rfonts.get(qn(att)) is not None:
            del rfonts.attrib[qn(att)]


def _estilos(doc):
    normal = doc.styles["Normal"]
    _fonte_estilo(normal, CORPO)
    pf = normal.paragraph_format
    pf.line_spacing = ESPACAMENTO
    pf.space_before, pf.space_after = Pt(0), DEPOIS
    pf.alignment = AL.JUSTIFY

    for nivel, tam in ((1, Pt(13)), (2, Pt(12)), (3, Pt(12))):
        st = doc.styles[f"Heading {nivel}"]
        _fonte_estilo(st, tam, negrito=nivel <= 2)
        hpf = st.paragraph_format
        hpf.line_spacing = ESPACAMENTO
        hpf.space_before = Pt(12 if nivel == 1 else 6)
        hpf.space_after = DEPOIS
        hpf.alignment = AL.LEFT
        hpf.keep_with_next = True

    for nome in ("List Bullet", "List Number", "Title", "Subtitle"):
        if nome in [s.name for s in doc.styles]:
            _fonte_estilo(doc.styles[nome], CORPO)
    lb = doc.styles["List Bullet"].paragraph_format
    lb.line_spacing, lb.space_after, lb.alignment = ESPACAMENTO, Pt(3), AL.JUSTIFY

    for nivel in (1, 2, 3):
        nome = f"TOC {nivel}"
        try:
            st = doc.styles[nome]
        except KeyError:
            st = doc.styles.add_style(nome, 1)
        _fonte_estilo(st, CORPO, negrito=nivel <= 2)
        st.paragraph_format.left_indent = Cm(0.5 * (nivel - 1))
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.line_spacing = 1.15


def _campo(paragrafo, instrucao, texto_provisorio=""):
    run = paragrafo.add_run()
    ini = OxmlElement("w:fldChar")
    ini.set(qn("w:fldCharType"), "begin")
    ins = OxmlElement("w:instrText")
    ins.set(qn("xml:space"), "preserve")
    ins.text = instrucao
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.extend([ini, ins, sep])
    if texto_provisorio:
        r2 = paragrafo.add_run(texto_provisorio)
        r2.font.name, r2.font.size = FONTE, CORPO
    fim_run = paragrafo.add_run()
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    fim_run._r.append(fim)
    return run


def _numero_de_pagina(sec):
    rodape = sec.footer
    p = rodape.paragraphs[0] if rodape.paragraphs else rodape.add_paragraph()
    p.alignment = AL.RIGHT
    p.paragraph_format.space_after = Pt(0)
    _campo(p, "PAGE")
    for r in p.runs:
        r.font.name, r.font.size = FONTE, CORPO
    # primeira pagina (capa) sem numero: rodape proprio vazio
    sec.first_page_footer.is_linked_to_previous = False


# --------------------------------------------------------------------------
# texto com marcacao
# --------------------------------------------------------------------------

_TOKENS = re.compile(
    r"(\*\*.+?\*\*|(?<![\*\w])\*(?!\s)[^*]+?(?<!\s)\*(?![\*\w])"
    r"|<sub>.*?</sub>|<sup>.*?</sup>)")


def escrever(p, texto, negrito=False, italico=False, tamanho=None):
    """Acrescenta ao paragrafo o texto, interpretando **negrito**, *italico*,
    <sub> e <sup>."""
    tam = Pt(tamanho) if isinstance(tamanho, (int, float)) else (
        tamanho or CORPO)
    for parte in _TOKENS.split(texto):
        if not parte:
            continue
        b, i, sub, sup = negrito, italico, False, False
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            parte, b = parte[2:-2], True
        elif parte.startswith("<sub>"):
            parte, sub = parte[5:-6], True
        elif parte.startswith("<sup>"):
            parte, sup = parte[5:-6], True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            parte, i = parte[1:-1], not italico
        run = p.add_run(parte)
        run.font.name, run.font.size = FONTE, tam
        run.font.bold, run.font.italic = b, i
        run.font.color.rgb = PRETO
        if sub:
            run.font.subscript = True
        if sup:
            run.font.superscript = True
    return p


def texto_simples(texto):
    """Remove a marcacao, para contagens e verificacoes."""
    t = re.sub(r"</?su[bp]>", "", texto)
    t = t.replace("**", "")
    t = re.sub(r"(?<![\*\w])\*(?!\s)([^*]+?)(?<!\s)\*(?![\*\w])", r"\1", t)
    return t


# --------------------------------------------------------------------------
# blocos elementares
# --------------------------------------------------------------------------


def titulo(doc, texto, nivel, nova_pagina=False):
    p = doc.add_heading("", level=nivel)
    escrever(p, texto.upper() if nivel == 1 else texto,
             negrito=nivel <= 2, tamanho=13 if nivel == 1 else 12)
    if nova_pagina:
        p.paragraph_format.page_break_before = True
    return p


def titulo_fora_do_indice(doc, texto, nova_pagina=True):
    """Titulo com o aspecto do nivel 1, mas que nao entra no indice."""
    p = doc.add_paragraph()
    p.alignment = AL.LEFT
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.keep_with_next = Pt(12), DEPOIS, True
    pf.page_break_before = nova_pagina
    escrever(p, texto.upper(), negrito=True, tamanho=13)
    return p


def subtitulo_formulario(doc, texto):
    """Seccao de um instrumento nos apendices (negrito, fora do indice)."""
    p = doc.add_paragraph()
    p.alignment = AL.LEFT
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.keep_with_next = Pt(10), Pt(4), True
    pf.line_spacing = 1.15
    escrever(p, texto, negrito=True)
    return p


def nova_pagina(paragrafo_obj):
    paragrafo_obj.paragraph_format.page_break_before = True
    return paragrafo_obj


def paragrafo(doc, texto, alinhamento=AL.JUSTIFY, negrito=False,
              italico=False, espaco_depois=None, espacamento=None,
              tamanho=None, recuo=None):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    if espaco_depois is not None:
        p.paragraph_format.space_after = Pt(espaco_depois)
    if espacamento is not None:
        p.paragraph_format.line_spacing = espacamento
    if recuo is not None:
        p.paragraph_format.left_indent = Cm(recuo)
    escrever(p, texto, negrito, italico, tamanho)
    return p


def lista(doc, itens, numerada=False):
    for n, item in enumerate(itens, 1):
        if numerada:
            p = doc.add_paragraph()
            p.alignment = AL.JUSTIFY
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Cm(0.9), Cm(-0.6)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Cm(0.9))
            escrever(p, f"{n}.\t{item}")
        else:
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = AL.JUSTIFY
            escrever(p, item)


def formula(doc, texto):
    p = doc.add_paragraph()
    p.alignment = AL.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    escrever(p, texto, italico=True)
    return p


def nota(doc, texto, alinhamento=AL.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(2), Pt(8), 1.0
    escrever(p, texto, tamanho=TAM_LEGENDA)
    return p


def campo_formulario(doc, texto):
    p = doc.add_paragraph()
    p.alignment = AL.LEFT
    pf = p.paragraph_format
    pf.space_after, pf.line_spacing = Pt(4), 1.15
    escrever(p, texto)
    return p


def quebra_de_pagina(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def legenda(doc, rotulo, numero, texto):
    """'Tabela 1: titulo', sempre acima do elemento."""
    p = doc.add_paragraph()
    p.alignment = AL.LEFT
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(8), Pt(3), 1.0
    pf.keep_with_next = True
    escrever(p, f"{rotulo} {numero}: {texto}", tamanho=TAM_LEGENDA)
    return p


def fonte_elemento(doc, texto):
    p = doc.add_paragraph()
    p.alignment = AL.LEFT
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(3), Pt(10), 1.0
    escrever(p, texto, tamanho=TAM_LEGENDA)
    return p


# --------------------------------------------------------------------------
# tabelas de tres linhas
# --------------------------------------------------------------------------


def _sem_bordas(celula):
    tcPr = celula._tc.get_or_add_tcPr()
    for antigo in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(antigo)
    bordas = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "nil")
        bordas.append(el)
    tcPr.append(bordas)


def _borda(celula, lado, tamanho=6):
    tcPr = celula._tc.get_or_add_tcPr()
    bordas = tcPr.find(qn("w:tcBorders"))
    if bordas is None:
        bordas = OxmlElement("w:tcBorders")
        tcPr.append(bordas)
    antigo = bordas.find(qn(f"w:{lado}"))
    if antigo is not None:
        bordas.remove(antigo)
    el = OxmlElement(f"w:{lado}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(tamanho))
    el.set(qn("w:color"), "000000")
    bordas.append(el)


def _repetir_cabecalho(linha):
    trPr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _nao_partir(linha):
    trPr = linha._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _escrever_celula(celula, texto, negrito, alinhamento, tamanho):
    celula.text = ""
    p = celula.paragraphs[0]
    p.alignment = alinhamento
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(1), Pt(1), 1.0
    pf.first_line_indent = Cm(0)
    escrever(p, str(texto), negrito=negrito, tamanho=tamanho)


def _fixar_grelha(t, larguras_cm, margem_cm=None):
    """Larguras fixas: tblLayout fixed, gridCol e tcW coerentes."""
    tbl = t._tbl
    tblPr = tbl.tblPr
    lay = tblPr.find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout")
        tblPr.append(lay)
    lay.set(qn("w:type"), "fixed")
    tw = OxmlElement("w:tblW")
    for antigo in tblPr.findall(qn("w:tblW")):
        tblPr.remove(antigo)
    tw.set(qn("w:w"), str(int(sum(larguras_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)
    if margem_cm is not None:
        mar = OxmlElement("w:tblCellMar")
        for lado in ("left", "right"):
            el = OxmlElement(f"w:{lado}")
            el.set(qn("w:w"), str(int(margem_cm * 567)))
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tblPr.append(mar)
    grid = tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    for gc, w in zip(cols, larguras_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for linha in t.rows:
        for j, celula in enumerate(linha._tr.findall(qn("w:tc"))):
            pass
        for j, w in enumerate(larguras_cm):
            if j < len(linha.cells):
                linha.cells[j].width = Cm(w)


def _larguras(t, larguras, n):
    if not larguras:
        larguras = [1] * n
    total = float(sum(larguras))
    _fixar_grelha(t, [LARGURA_UTIL_CM * w / total for w in larguras])


def _alinhamento_coluna(linhas, j):
    """Numeros a' direita/centro; texto a' esquerda."""
    valores = [str(l[j]).replace("**", "").strip() for l in linhas
               if j < len(l)]
    if valores and all(re.fullmatch(r"[\d.,%()\s+<>=-]*", v or "0")
                       for v in valores):
        return AL.CENTER
    return AL.LEFT


def tabela(doc, cabecalho, linhas, larguras=None, tamanho=TAM_TABELA,
           linhas_de_acento=(), negrito_ultima=False):
    n = len(cabecalho)
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    alinh = [_alinhamento_coluna(linhas, j) for j in range(n)]
    for j, texto in enumerate(cabecalho):
        _escrever_celula(t.rows[0].cells[j], texto, True,
                         AL.CENTER if alinh[j] == AL.CENTER else AL.LEFT,
                         tamanho)
    _repetir_cabecalho(t.rows[0])
    for k, linha in enumerate(linhas):
        celulas = t.add_row().cells
        ultima = negrito_ultima and k == len(linhas) - 1
        for j in range(n):
            texto = linha[j] if j < len(linha) else ""
            _escrever_celula(celulas[j], texto, ultima, alinh[j], tamanho)
    for linha in t.rows:
        _nao_partir(linha)
        for celula in linha.cells:
            _sem_bordas(celula)
    for celula in t.rows[0].cells:
        _borda(celula, "top", 8)
        _borda(celula, "bottom", 6)
    for celula in t.rows[-1].cells:
        _borda(celula, "bottom", 8)
    for idx in linhas_de_acento:
        for celula in t.rows[idx + 1].cells:
            _borda(celula, "top", 4)
    _larguras(t, larguras, n)
    return t


def tabela_cronograma(doc, meses, actividades, tamanho=TAM_TABELA):
    """meses: lista de (ano, abreviatura); actividades: (texto, set de
    indices 0..n-1). Cabecalho em duas linhas (ano agrupado)."""
    n = len(meses)
    t = doc.add_table(rows=2, cols=n + 1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    anos = []
    for j, (ano, _) in enumerate(meses):
        if not anos or anos[-1][0] != ano:
            anos.append([ano, j, j])
        else:
            anos[-1][2] = j
    c0 = t.cell(0, 0).merge(t.cell(1, 0))
    _escrever_celula(c0, "Actividade", True, AL.LEFT, tamanho)
    for ano, a, b in anos:
        c = t.cell(0, a + 1)
        if b > a:
            c = c.merge(t.cell(0, b + 1))
        _escrever_celula(c, str(ano), True, AL.CENTER, tamanho)
    for j, (_, mes) in enumerate(meses):
        _escrever_celula(t.cell(1, j + 1), mes, True, AL.CENTER, tamanho - 1)
    _repetir_cabecalho(t.rows[0])
    _repetir_cabecalho(t.rows[1])
    for texto, idx in actividades:
        celulas = t.add_row().cells
        _escrever_celula(celulas[0], texto, False, AL.LEFT, tamanho)
        for j in range(n):
            _escrever_celula(celulas[j + 1], "X" if j in idx else "", False,
                             AL.CENTER, tamanho)
    for linha in t.rows:
        _nao_partir(linha)
        for celula in linha.cells:
            _sem_bordas(celula)
    for celula in t.rows[0].cells:
        _borda(celula, "top", 8)
    for celula in t.rows[1].cells:
        _borda(celula, "bottom", 6)
    for celula in t.rows[-1].cells:
        _borda(celula, "bottom", 8)
    larg_mes = min(0.85, 9.6 / n)
    larg_act = LARGURA_UTIL_CM - larg_mes * n
    _fixar_grelha(t, [larg_act] + [larg_mes] * n, margem_cm=0.06)
    return t


# --------------------------------------------------------------------------
# pre-textuais
# --------------------------------------------------------------------------


def linhas_centradas(doc, linhas, negrito=False, espaco=2, tamanho=None,
                     espacamento=1.15):
    ps = []
    for texto in linhas:
        p = paragrafo(doc, texto, AL.CENTER, negrito=negrito,
                      espaco_depois=espaco, tamanho=tamanho)
        p.paragraph_format.line_spacing = espacamento
        ps.append(p)
    return ps


def vazios(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)


def paragrafo_direita(doc, texto, recuo_cm=8.0):
    """Nota de apresentacao da folha de rosto, recuada a' direita."""
    p = doc.add_paragraph()
    p.alignment = AL.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent, pf.line_spacing, pf.space_after = Cm(recuo_cm), 1.0, Pt(6)
    escrever(p, texto)
    return p


def sigla(doc, abreviatura, significado):
    p = doc.add_paragraph()
    p.alignment = AL.LEFT
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Cm(3.0), Cm(-3.0)
    pf.space_after, pf.line_spacing = Pt(2), 1.15
    pf.tab_stops.add_tab_stop(Cm(3.0), WD_TAB_ALIGNMENT.LEFT)
    escrever(p, f"{abreviatura}\t{significado}")
    return p


def indice(doc):
    p = doc.add_paragraph()
    _campo(p, r'TOC \o "1-3" \h \z \u',
           "Índice por actualizar: no Word, clicar sobre esta área e "
           "premir F9 (Actualizar o índice inteiro).")
    return p


def referencia(doc, numero, texto):
    p = doc.add_paragraph()
    p.alignment = AL.LEFT
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Cm(0.75), Cm(-0.75)
    pf.line_spacing, pf.space_after = 1.0, Pt(6)
    pf.tab_stops.add_tab_stop(Cm(0.75))
    escrever(p, f"{numero}.\t{texto}")
    return p


def imagem(doc, caminho, largura_cm=LARGURA_UTIL_CM):
    p = doc.add_paragraph()
    p.alignment = AL.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(caminho, width=Cm(largura_cm))
    return p
